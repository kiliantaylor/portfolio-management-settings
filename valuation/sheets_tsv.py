"""One-tab Google Sheets version of a slim model (sheets_slim.py output), written as TSV so the Drive connector can upload it as
plain text (Drive converts text/tab-separated-values to a Google Sheet and evaluates cells that start with '=').
The five tabs are stacked on one sheet in this order with a header row each; every formula is rewritten to the stacked row
numbers and defined names are inlined. Usage: python3 sheets_tsv.py DATE SYM... -> runs/DATE/SYM/sheet_SYM.tsv
"""
import os, re, sys, openpyxl
from openpyxl.utils import column_index_from_string
HERE = os.path.dirname(os.path.abspath(__file__))
ORDER = ['Summary', 'Scenarios', 'Methods', 'Distribution', 'Assumptions']
REF = re.compile(r"(?:(?P<sh>[A-Za-z_]+)!)?(?<![A-Za-z_\.\d])(?P<ca>\$?)(?P<c>[A-Z]{1,3})(?P<ra>\$?)(?P<r>\d+)(?![\d\(A-Za-z_])")
RANGE = re.compile(r"([A-Za-z_]+)!(\$?[A-Z]{1,3}\$?\d+):(\$?[A-Z]{1,3}\$?\d+)")
date, syms = sys.argv[1], sys.argv[2:]
for s in syms:
    wb = openpyxl.load_workbook(os.path.join(HERE, 'runs', date, s, f'sheet_{s}.xlsx'))
    off, row0 = {}, 0
    for n in ORDER: off[n] = row0 + 1; row0 += wb[n].max_row + 2  # one title row per block, one blank row after
    names = {n: dn.attr_text.replace('$', '') for n, dn in wb.defined_names.items()}
    grid = {}
    for n in ORDER:
        ws = wb[n]; grid[(off[n], 1)] = f'{n.upper()}'
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None: continue
                if isinstance(v, str) and v.startswith('='):
                    for nm, t in names.items():
                        sh_, a1 = t.split('!'); m_ = re.match(r'([A-Z]+)(\d+)$', a1)
                        v = re.sub(rf'(?<![A-Za-z_!"]){re.escape(nm)}(?![A-Za-z_\d])', f'{sh_}!${m_.group(1)}${m_.group(2)}', v)
                    v = RANGE.sub(lambda m: f"{m.group(1)}!{m.group(2)}:{m.group(1)}!{m.group(3)}", v)  # sheet prefix on both ends
                    parts = re.split(r'("[^"]*")', v); out = []
                    for p in parts:
                        if p.startswith('"'): out.append(p); continue
                        def sub(m):
                            sh = m.group('sh') or n; sh = sh if sh in off else n
                            return f"{m.group('ca')}{m.group('c')}{m.group('ra')}{int(m.group('r')) + off[sh]}"
                        out.append(REF.sub(sub, p))
                    v = ''.join(out).replace('$$', '$')
                elif isinstance(v, str) and v[:1] in '=+-@': v = "'" + v
                if n == 'Assumptions' and c.row > 4 and c.column == 8: continue  # 'All sources' index; full list stays in the Excel model
                if n == 'Assumptions' and c.row > 4 and c.column == 4 and ws.cell(c.row, 2).value and ' · ' in str(ws.cell(c.row, 2).value): continue  # agents' own reasoning: Excel model only
                grid[(c.row + off[n], c.column)] = v
    R = max(r for r, _ in grid); C = max(c for _, c in grid)
    def cell(v):
        if v is None: return ''
        v = str(v)
        return '"' + v.replace('"', '""') + '"' if any(ch in v for ch in '\t\n"') else v
    grid = {k: (v.replace('Blue = input you can change, black = formula, green = link to another sheet. Yellow = key assumption.', 'Change any number in the Value column and the rest recalculates; column C tracks the live price. Agents\' reasoning and every source link are in the Excel model on the Price Targets page.') if isinstance(v, str) else v) for k, v in grid.items()}
    lines = ['\t'.join(cell(grid.get((r, c))) for c in range(1, C + 1)).rstrip('\t') for r in range(1, R + 1)]
    out = os.path.join(HERE, 'runs', date, s, f'sheet_{s}.tsv'); open(out, 'w').write('\n'.join(lines) + '\n')
    print(s, os.path.getsize(out), 'rows', R, 'offsets', off)
