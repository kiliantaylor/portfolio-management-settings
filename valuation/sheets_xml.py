"""Write a slim model (sheets_slim.py output) as Excel 2003 XML Spreadsheet text, which the Google Drive connector can upload
as plain text and convert to a multi-tab Google Sheet with working formulas (binary .xlsx is too large to send inline).
Usage: python3 sheets_xml.py DATE SYM... -> runs/DATE/SYM/sheet_SYM.xml (formulas converted to R1C1, defined names inlined).
Upload: Drive create_file, contentMimeType application/vnd.ms-excel, textContent = the file, parent 'Price target models'.
"""
import os, re, sys, openpyxl
from xml.sax.saxutils import escape
from openpyxl.utils import column_index_from_string
HERE = os.path.dirname(os.path.abspath(__file__))
REF = re.compile(r"(?<![A-Za-z_\.\d])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\d\(A-Za-z_])")
def r1c1(formula, row, col):
    out, parts = [], re.split(r'("[^"]*")', formula)
    for p in parts:
        if p.startswith('"'): out.append(p); continue
        def sub(m):
            ca, c, ra, r = m.group(1), column_index_from_string(m.group(2)), m.group(3), int(m.group(4))
            rs = f'R{r}' if ra else ('R' if r == row else f'R[{r-row}]')
            cs = f'C{c}' if ca else ('C' if c == col else f'C[{c-col}]')
            return rs + cs
        out.append(REF.sub(sub, p))
    return ''.join(out)
date, syms = sys.argv[1], sys.argv[2:]
for s in syms:
    wb = openpyxl.load_workbook(os.path.join(HERE, 'runs', date, s, f'sheet_{s}.xlsx'))
    names = {}
    for n, dn in wb.defined_names.items():
        names[n] = dn.attr_text.replace('$', '')  # e.g. Summary!B5 -> absolute below
    L = ['<?xml version="1.0"?><?mso-application progid="Excel.Sheet"?>',
         '<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">']
    for ws in wb.worksheets:
        L.append(f'<Worksheet ss:Name="{escape(ws.title)}"><Table>')
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            cells, last = [], 0
            for c in row:
                v = c.value
                if v is None: continue
                idx = f' ss:Index="{c.column}"' if c.column != last + 1 else ''; last = c.column
                if isinstance(v, str) and v.startswith('='):
                    f = v
                    for n, t in names.items():
                        sh, ref = t.split('!'); m = re.match(r'([A-Z]+)(\d+)', ref)
                        f = re.sub(rf'(?<![A-Za-z_!]){re.escape(n)}(?![A-Za-z_\d])', f"{sh}!${m.group(1)}${m.group(2)}", f)
                    cells.append(f'<Cell{idx} ss:Formula="{escape(r1c1(f, c.row, c.column), {chr(34): "&quot;"})}"/>')
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    cells.append(f'<Cell{idx}><Data ss:Type="Number">{v}</Data></Cell>')
                else:
                    cells.append(f'<Cell{idx}><Data ss:Type="String">{escape(str(v))}</Data></Cell>')
            L.append(f'<Row ss:Index="{row[0].row}">' + ''.join(cells) + '</Row>' if cells else '')
        L.append('</Table></Worksheet>')
    L.append('</Workbook>')
    out = os.path.join(HERE, 'runs', date, s, f'sheet_{s}.xml'); open(out, 'w').write('\n'.join(x for x in L if x))
    print(s, os.path.getsize(out), 'names', list(names))
