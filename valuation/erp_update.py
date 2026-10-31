"""Discount-rate inputs for the fixed DCF (card dcf_v2_reenable, Kilian chose "Turn it on, tested weights" 2026-10-21 08:05Z).
Writes data/erp/erp.json, which model_engine.py reads for every DCF leg:
  erp     Aswath Damodaran's implied equity risk premium for the S&P 500, column "ERP (T12m)" of his monthly file
          https://pages.stern.nyu.edu/~adamodar/pc/implprem/ERPbymonth.xlsx (sheet "Historical ERP", start of month). The same
          column was tested on the WRDS harness (wrds/code/dcf_damodaran_erp.py, results/dcf_bias/damodaran_erp.csv).
  rf      the 10-year US Treasury yield, last close of Yahoo's ^TNX chart feed.
Usage: python3 erp_update.py        (nightly; the file is re-downloaded when its month is older than the current month)"""
import json, os, sys, datetime as dt, subprocess, tempfile
HERE = os.path.dirname(os.path.abspath(__file__)); D = os.path.join(HERE, 'data', 'erp'); OUT = os.path.join(D, 'erp.json')
URL = 'https://pages.stern.nyu.edu/~adamodar/pc/implprem/ERPbymonth.xlsx'
TNX = 'https://query1.finance.yahoo.com/v8/finance/chart/%5ETNX?range=10d&interval=1d'

def fetch(url, path):
    subprocess.run(['curl', '-sS', '-L', '-f', '-A', 'Mozilla/5.0', '-o', path, url], check=True, timeout=60)

def read_erp(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)['Historical ERP']; hdr = [c.value for c in ws[1]]
    i = hdr.index('ERP (T12m)'); last = None
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is not None and isinstance(r[i], (int, float)): last = (r[0].date().isoformat(), float(r[i]), r[hdr.index('T.Bond Rate')])
    return last

def main():
    os.makedirs(D, exist_ok=True)
    cur = json.load(open(OUT)) if os.path.exists(OUT) else {}
    today = dt.date.today(); m0 = today.replace(day=1).isoformat()
    if cur.get('erp_month', '') < m0:
        tmp = tempfile.mkdtemp(); p = os.path.join(tmp, 'ERPbymonth.xlsx')
        try:
            fetch(URL, p); mon, erp, tb = read_erp(p)
            if mon > cur.get('erp_month', ''):
                import shutil; shutil.copyfile(p, os.path.join(D, f"ERPbymonth_{mon}.xlsx"))
                cur.update(erp=erp, erp_month=mon, erp_tbond_at_month_start=tb, erp_source=f'Damodaran, Implied ERP by month, ERP (T12m), {mon} ({URL})')
        except Exception as e: print('ERP download failed, keeping', cur.get('erp_month'), e, file=sys.stderr)
    try:
        out = subprocess.run(['curl', '-sS', '-f', '-A', 'Mozilla/5.0', TNX], capture_output=True, check=True, timeout=60).stdout
        r = json.loads(out)['chart']['result'][0]; pts = [(t, c) for t, c in zip(r['timestamp'], r['indicators']['quote'][0]['close']) if c]
        t, c = pts[-1]; cur.update(rf=round(c / 100, 5), rf_date=dt.datetime.fromtimestamp(t, dt.timezone.utc).date().isoformat(),
                                   rf_source='10-year US Treasury yield, Yahoo ^TNX daily close')
    except Exception as e: print('10y download failed, keeping', cur.get('rf_date'), e, file=sys.stderr)
    cur['updated'] = dt.datetime.now(dt.timezone.utc).strftime('%Y-%m-%dT%H:%MZ')
    json.dump(cur, open(OUT, 'w'), indent=1); print(json.dumps(cur))

if __name__ == '__main__': main()
