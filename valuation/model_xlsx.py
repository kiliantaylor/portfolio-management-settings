"""Excel model of one name's price-target round: the supervisor's assumption table (each agent's value, reasoning and
linked sources), every valuation method, scenarios, and the probability ladder with the blend as live formulas.

python3 model_xlsx.py DATE [SYM ...] [--jobs N] [--force]   ->  runs/DATE/SYM/model_SYM.xlsx
DATE is the re-mark date (runs/DATE/SYM/result.json); the agents' files come from result['research_as_of'].
Speed (2026-10-21, Kilian approved the Price Targets speedup list): workbooks are built N at a time (default 4, one
LibreOffice per worker), each in a local temp folder and then copied to runs/DATE/SYM, and a workbook whose inputs
(result.json, the research files it reads, this script and combine.py) are unchanged since it was last built for DATE is
kept (runs/DATE/SYM/model_SYM.inputs.md5); --force rebuilds anyway.
"""
from combine import survival_fn
import json, os, sys, glob
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__)).replace('/mnt/attach/project-files', '/mnt/project-files')
F = 'Arial'
BLUE, BLACK, GREEN, GREY = '0000FF', '000000', '008000', '6B6B6B'
HDR = PatternFill('solid', fgColor='1F2937'); SUB = PatternFill('solid', fgColor='E5E7EB')
KEY = PatternFill('solid', fgColor='FFFF00'); OUT = PatternFill('solid', fgColor='EAF3EA')
THIN = Border(bottom=Side(style='thin', color='D1D5DB'))
WRAP = Alignment(wrap_text=True, vertical='top')
USD, PCT, PROB, NUM = '$#,##0.00;($#,##0.00);-', '0.0%;(0.0%);-', '0.000', '#,##0.00;(#,##0.00);-'


def jl(p, default=None):
    try:
        return json.load(open(p))
    except Exception:
        return default


def cell(ws, r, c, v, color=BLACK, bold=False, fmt=None, fill=None, wrap=False, size=10, link=None, note=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font = Font(name=F, size=size, color=color, bold=bold, underline='single' if link else None)
    if fmt: x.number_format = fmt
    if fill: x.fill = fill
    if wrap: x.alignment = WRAP
    else: x.alignment = Alignment(vertical='top')
    if link and link.startswith('#'):
        x.hyperlink = Hyperlink(ref=x.coordinate, location=link[1:])
    elif link:
        x.hyperlink = link
    if note: x.comment = Comment(note, 'model')
    return x


def header(ws, r, cols, widths=None):
    for i, h in enumerate(cols, 1):
        x = cell(ws, r, i, h, color='FFFFFF', bold=True, fill=HDR, wrap=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[CL(i)].width = w


def title(ws, text, sub):
    cell(ws, 1, 1, text, bold=True, size=14)
    cell(ws, 2, 1, sub, color=GREY)


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def short(t, n=60):
    t = (t or '').strip()
    return t if len(t) <= n else t[:n - 1] + '…'


def build(date, sym, work=None):
    rdir = os.path.join(HERE, 'runs', date, sym)
    res = jl(os.path.join(rdir, 'result.json'))
    if not res:
        print(sym, 'no result.json'); return None
    rdate = res.get('research_as_of') or date
    adir = os.path.join(HERE, 'runs', rdate, sym)
    agents = [a for a in (jl(p) for p in sorted(glob.glob(os.path.join(adir, 'agent_*.json')))) if a]
    r1 = {os.path.basename(p)[:-5]: jl(p) for p in glob.glob(os.path.join(adir, 'round1', 'agent_*.json'))}
    sup = jl(os.path.join(adir, 'supervisor.json'), {}) or {}
    sup1 = jl(os.path.join(adir, 'round1', 'supervisor.json'), {}) or {}
    ma = jl(os.path.join(adir, 'model_assumptions.json'), {}) or {}
    names = [a.get('agent_id', f'agent_{i+1}') for i, a in enumerate(agents)]
    nA = len(agents)
    S = res['spot']; m = res['method']; an = res.get('analysts') or {}

    wb = Workbook()
    # ---------------- Sources (built first so other sheets can link to rows) ----------------
    src = wb.active; src.title = 'Sources'
    srows = {}  # (agent, idx) -> (row, url, title)
    title(src, f'{sym} sources', 'Every source the forecasters and the supervisor cited. IDs: A1.3 = agent 1, source 3; S.2 = supervisor, source 2; X.4 = found by the supervisor for an assumption the agent left unsourced.')
    header(src, 4, ['ID', 'Cited by', 'Title', 'Date', 'Link'], [8, 12, 70, 12, 60])
    r = 5
    added, seen = [], set()  # sources the supervisor found for assumptions the agents left unsourced
    for a_ in ma.get('assumptions') or []:
        for ag in (a_.get('agents') or {}).values():
            for x in (ag or {}).get('added_sources') or []:
                if isinstance(x, dict) and x.get('url') and x['url'] not in seen:
                    seen.add(x['url']); added.append(x)
    for who, lst, tag in [(n, a.get('sources') or [], f'A{i+1}') for i, (n, a) in enumerate(zip(names, agents))] + [('supervisor', sup.get('sources') or [], 'S'), ('supervisor (added)', added, 'X')]:
        for j, s in enumerate(lst, 1):
            s = s if isinstance(s, dict) else {'title': str(s), 'url': str(s) if str(s).startswith('http') else ''}
            url = s.get('url') or ''
            cell(src, r, 1, f'{tag}.{j}'); cell(src, r, 2, who); cell(src, r, 3, s.get('title') or url, wrap=True)
            cell(src, r, 4, s.get('date') or '')
            cell(src, r, 5, url, color=BLUE if url.startswith('http') else BLACK, link=url if url.startswith('http') else None)
            srows[(who, j)] = (r, url, s.get('title') or url)
            r += 1
    src.freeze_panes = 'A5'

    # ---------------- Summary ----------------
    sm = wb.create_sheet('Summary', 0)
    title(sm, f'{sym} 12-month price target model',
          f'Close {res["as_of"]} · agents\' research {rdate} · horizon {res["horizon"]}. Blue = input you can change, black = formula, '
          f'green = link to another sheet. Yellow = key assumption.')
    sm.column_dimensions['A'].width = 38; sm.column_dimensions['B'].width = 16; sm.column_dimensions['C'].width = 16
    sm.column_dimensions['D'].width = 70
    cell(sm, 4, 1, 'Inputs', bold=True, fill=SUB); cell(sm, 4, 2, 'Value', bold=True, fill=SUB)
    cell(sm, 4, 3, '', fill=SUB); cell(sm, 4, 4, 'Where it comes from', bold=True, fill=SUB)
    special = bool(m.get('special_situation'))
    iv = res.get('iv30') or res.get('atm_iv_1y')
    inputs = [
        ('Spot', 'Last close', S, USD, f'Close on {res["as_of"]}'),
        ('W_Mkt', 'Weight: options market', m['w_market'], PCT, 'AIA Forecaster default (market 67%); flips to 33% when the option chain is thin'),
        ('W_Agents', 'Weight: agents', m['w_llm'], PCT, 'AIA Forecaster default (agents 33%)'),
        ('W_Analyst', 'Weight: analyst consensus', an.get('weight', 0) or 0, PCT, 'Taken pro rata from the two above; 0 when there are no targets or in a special situation'),
        ('Analyst_Median', 'Analysts\' median target', an.get('median') or 0, USD, f'{an.get("n", "")} analysts, {an.get("source", "")} as of {an.get("as_of", "")}'),
        ('Sup_Used', 'Use supervisor ladder (1 = yes)', 1 if res.get('supervisor_used') else 0, '0', f'Supervisor confidence: {res.get("supervisor_confidence")}; its ladder replaces the agents\' mean only when high'),
        ('Platt_Shift', 'Platt log shift of the agents\' curve', (res.get('platt') or {}).get('log_shift', 0) or 0, '0.0000',
         f'combine.py: P(close >= spot) {res["platt"]["p_up_raw"]:.3f} Platt-scaled (a = {m["platt_a"]:.3f}) to {res["platt"]["p_up_platt"]:.3f}, '
         'then the whole agents\' curve is shifted in log price by this amount' if res.get('platt') else 'no Platt step'),
        ('Thesis_Break', 'Thesis-break price (supervisor)', res['thesis_break'].get('reconciled') or 0, USD, 'Supervisor\'s reconciled thesis break (agents\' values on the Reasoning sheet)'),
        ('IV30', '30-day implied vol', iv or 0, PCT, f'Cboe model-free 30D IV as of {res.get("iv30_date")}; sets the stop noise floor (2 one-month moves)'),
    ]
    r = 5
    for nm, lab, v, fmt, why in inputs:
        cell(sm, r, 1, lab); cell(sm, r, 2, v, color=BLUE, fmt=fmt, fill=KEY if nm in ('W_Mkt', 'W_Agents', 'W_Analyst', 'Analyst_Median') else None)
        cell(sm, r, 4, why, color=GREY, wrap=True)
        wb.defined_names[nm] = DefinedName(nm, attr_text=f"Summary!$B${r}")
        r += 1

    # ---------------- Distribution ----------------
    ds = wb.create_sheet('Distribution')
    lad = res['ladder']; n = len(lad); r0 = 6; r1_ = r0 + n - 1
    title(ds, f'{sym} probability of closing at or above each price on {res["horizon"]}',
          'Each column is P(close >= K). Agents\' curve is Platt-shifted, blended with the options market, then analysts at W_Analyst (0 since 2026-10-05). '
          'Interpolation is linear in logit(p) vs log(K), as in combine.py.')
    ac = list(range(4, 4 + nA))  # agent columns
    c_mean, c_sup, c_used, c_lu, c_xp, c_jp, c_pl, c_lm, c_xa, c_ja, c_an, c_ex, c_fin, c_lf, c_pub = range(4 + nA, 4 + nA + 15)
    heads = ['Price K', 'ln K', 'Options market'] + [f'Agent {i+1}' for i in range(nA)] + [
        'Agents\' mean', 'Supervisor', 'Agents\' curve used', 'logit(used)', 'K shifted', 'seg', 'Agents Platt-shifted',
        'logit(market)', 'K re-centred', 'seg', 'Analyst curve', 'Blend ex analysts', 'Final blend', 'logit(final)', 'Published (combine.py)']
    header(ds, r0 - 1, heads, [11, 8, 11] + [10] * nA + [11, 11, 11, 9, 10, 6, 12, 9, 10, 6, 11, 11, 11, 9, 12])
    KR = f'Distribution!$A${r0}:$A${r1_}'; LNR = f'Distribution!$B${r0}:$B${r1_}'
    rng = lambda c: f'Distribution!${CL(c)}${r0}:${CL(c)}${r1_}'
    clip = lambda x: f'MIN(MAX({x},0.0001),0.9999)'
    logit = lambda x: f'LN({clip(x)}/(1-{clip(x)}))'
    sup_lad = {round(l['K']): l['p'] for l in (sup.get('ladder') or [])}

    def agent_p(a, K):
        L = a.get('ladder') or []
        if not L: return None
        # the agents' ladder is on the research day's strikes; re-marks interpolate it onto today's, as combine.py does
        L = sorted(L, key=lambda l: l['K'])
        return round(survival_fn([l['K'] for l in L], [min(max(l['p'], 1e-4), 1 - 1e-4) for l in L])(K), 4)

    def interp(xc, jc, lc):
        i = lambda c, o='': f'INDEX({rng(c)},{jc}{o})'
        return (f'1/(1+EXP(-({i(lc)}+(LN({xc})-INDEX({LNR},{jc}))/(INDEX({LNR},{jc}+1)-INDEX({LNR},{jc}))'
                f'*(INDEX({rng(lc)},{jc}+1)-INDEX({rng(lc)},{jc})))))')

    for k, l in enumerate(lad):
        rr = r0 + k; K = l['K']
        cell(ds, rr, 1, round(K, 2), color=BLUE, fmt=USD)
        cell(ds, rr, 2, f'=LN(A{rr})', fmt='0.000')
        cell(ds, rr, 3, l['market'], color=BLUE, fmt=PROB)
        for i, a in enumerate(agents):
            p = agent_p(a, K)
            cell(ds, rr, ac[i], p if p is not None else '', color=BLUE, fmt=PROB)
        cell(ds, rr, c_mean, f'=AVERAGE({CL(ac[0])}{rr}:{CL(ac[-1])}{rr})', fmt=PROB)
        sp = agent_p(sup, K) if sup.get('ladder') else None
        cell(ds, rr, c_sup, sp if sp is not None else '', color=BLUE, fmt=PROB)
        cell(ds, rr, c_used, f'=IF(AND(Sup_Used=1,ISNUMBER({CL(c_sup)}{rr})),{CL(c_sup)}{rr},{CL(c_mean)}{rr})', fmt=PROB)
        cell(ds, rr, c_lu, '=' + logit(f'{CL(c_used)}{rr}'), fmt='0.000')
        cell(ds, rr, c_xp, f'=A{rr}*EXP(-Platt_Shift)', fmt=USD)
        cell(ds, rr, c_jp, f'=MIN(MAX(IFERROR(MATCH({CL(c_xp)}{rr},{KR},1),1),1),{n-1})', fmt='0')
        cell(ds, rr, c_pl, '=' + interp(f'{CL(c_xp)}{rr}', f'{CL(c_jp)}{rr}', c_lu), fmt=PROB)
        cell(ds, rr, c_lm, '=' + logit(f'C{rr}'), fmt='0.000')
        cell(ds, rr, c_xa, f'=IF(Analyst_Median>0,A{rr}*Market_Median/Analyst_Median,A{rr})', fmt=USD)
        cell(ds, rr, c_ja, f'=MIN(MAX(IFERROR(MATCH({CL(c_xa)}{rr},{KR},1),1),1),{n-1})', fmt='0')
        cell(ds, rr, c_an, '=' + interp(f'{CL(c_xa)}{rr}', f'{CL(c_ja)}{rr}', c_lm), fmt=PROB)
        cell(ds, rr, c_ex, f'=W_Mkt*C{rr}+W_Agents*{CL(c_pl)}{rr}', fmt=PROB)
        cell(ds, rr, c_fin, f'=(1-W_Analyst)*{CL(c_ex)}{rr}+W_Analyst*{CL(c_an)}{rr}', fmt=PROB, fill=OUT)
        cell(ds, rr, c_lf, '=' + logit(f'{CL(c_fin)}{rr}'), fmt='0.000')
        cell(ds, rr, c_pub, l['final'], color=BLUE, fmt=PROB)
    ds.freeze_panes = ds.cell(row=r0, column=2)
    cell(ds, r1_ + 2, 1, 'Agents\' curves are their final-round answers (Reasoning sheet shows round 1). Market and published columns '
         'come from combine.py; the model reproduces the final blend within rounding (combine.py also enforces a non-increasing curve).',
         color=GREY)

    def q_formula(c, s):
        """Price where column c (descending P) crosses survival level s, logit-in-log-K interpolation."""
        lc = c_lf if c == c_fin else c_lm
        j = f'MIN(MAX(IFERROR(MATCH({s},{rng(c)},-1),1),1),{n-1})'
        Li = lambda o: f'INDEX({rng(lc)},{j}{o})'
        Ki = lambda o: f'INDEX({LNR},{j}{o})'
        return f'=EXP({Ki("")}+({logit(str(s))}-{Li("")})/({Li("+1")}-{Li("")})*({Ki("+1")}-{Ki("")}))'

    # market median (for re-centring analysts) as a defined name on Summary
    r += 1
    cell(sm, r, 1, 'Options-market median', ); cell(sm, r, 2, q_formula(3, 0.5), color=GREEN, fmt=USD)
    cell(sm, r, 4, 'Median of the options-implied distribution (Distribution sheet); the analyst curve is this curve re-centred on the analysts\' median', color=GREY, wrap=True)
    wb.defined_names['Market_Median'] = DefinedName('Market_Median', attr_text=f'Summary!$B${r}')
    r += 2

    # ---------------- Methods ----------------
    me = wb.create_sheet('Methods')
    title(me, f'{sym} valuation methods by agent', 'Each agent\'s methods, the value per share and the weight it gave each. '
          'Blended value = weight-averaged methods; stated value = what the agent reported (it may round or add a deduction).')
    header(me, 4, ['Agent', 'Method', 'Value / share', 'Weight', 'Value x weight', 'Calculation (assumption IDs on the Assumptions sheet)', 'Inputs and reasoning'],
           [10, 20, 13, 9, 13, 44, 90])
    mrow = 5; agent_blend = {}
    ma_methods = {}
    for mm in ma.get('methods') or []:
        ma_methods.setdefault(mm.get('agent'), []).append(mm)
    for i, (nm, a) in enumerate(zip(names, agents)):
        ms = (a.get('fair_value') or {}).get('methods') or []
        first = mrow
        for mth in ms:
            extra = next((x for x in ma_methods.get(nm, []) if x.get('name') == mth.get('name')), {})
            cell(me, mrow, 1, f'Agent {i+1}'); cell(me, mrow, 2, mth.get('name'), wrap=True)
            cell(me, mrow, 3, num(mth.get('value')), color=BLUE, fmt=USD)
            cell(me, mrow, 4, num(mth.get('weight')) or 0, color=BLUE, fmt=PCT)
            cell(me, mrow, 5, f'=C{mrow}*D{mrow}', fmt=USD)
            calc = extra.get('formula') or ''
            if extra.get('assumption_ids'): calc += f'  [{", ".join(extra["assumption_ids"])}]'
            cell(me, mrow, 6, calc, wrap=True)
            cell(me, mrow, 7, mth.get('inputs') or mth.get('assumptions') or extra.get('reasoning') or '', wrap=True)
            mrow += 1
        last = mrow - 1
        cell(me, mrow, 1, f'Agent {i+1}', bold=True); cell(me, mrow, 2, 'Blended value', bold=True)
        cell(me, mrow, 3, f'=IF(SUM(D{first}:D{last})>0,SUMPRODUCT(C{first}:C{last},D{first}:D{last})/SUM(D{first}:D{last}),"")' if ms else '', fmt=USD, bold=True, fill=OUT)
        cell(me, mrow, 4, f'=SUM(D{first}:D{last})' if ms else '', fmt=PCT)
        cell(me, mrow, 6, 'Stated research target (12m)', bold=True)
        cell(me, mrow, 5, num((a.get('fair_value') or {}).get('value')), color=BLUE, fmt=USD, bold=True)
        cell(me, mrow, 7, (a.get('fair_value') or {}).get('method') or '', wrap=True, color=GREY)
        for c in range(1, 8): me.cell(row=mrow, column=c).border = THIN
        agent_blend[nm] = (mrow, first, last)
        mrow += 2
    me.freeze_panes = 'C5'

    # Summary: research targets (12m)
    cell(sm, r, 1, 'Research target (12m) per share', bold=True, fill=SUB); cell(sm, r, 2, 'Method blend', bold=True, fill=SUB)
    cell(sm, r, 3, 'Stated', bold=True, fill=SUB); cell(sm, r, 4, 'Note', bold=True, fill=SUB); r += 1
    fv_first = r
    for i, nm in enumerate(names):
        br = agent_blend[nm][0]
        cell(sm, r, 1, f'Agent {i+1}'); cell(sm, r, 2, f'=Methods!C{br}', color=GREEN, fmt=USD)
        cell(sm, r, 3, f'=Methods!E{br}', color=GREEN, fmt=USD); r += 1
    fv_last = r - 1
    cell(sm, r, 1, 'Agents\' average', bold=True); cell(sm, r, 2, f'=AVERAGE(B{fv_first}:B{fv_last})', fmt=USD, bold=True)
    cell(sm, r, 3, f'=AVERAGE(C{fv_first}:C{fv_last})', fmt=USD, bold=True); r += 1
    cell(sm, r, 1, 'Agents\' median'); cell(sm, r, 2, f'=MEDIAN(B{fv_first}:B{fv_last})', fmt=USD)
    cell(sm, r, 3, f'=MEDIAN(C{fv_first}:C{fv_last})', fmt=USD); r += 1
    sfv = sup.get('fair_value') if isinstance(sup.get('fair_value'), dict) else {}
    cell(sm, r, 1, 'Supervisor\'s reconciled research target (12m)', bold=True)
    cell(sm, r, 3, num(sfv.get('value')) or res.get('fair_value_median'), color=BLUE, fmt=USD, bold=True, fill=KEY)
    cell(sm, r, 4, sfv.get('reason') or res.get('fair_value_reason') or '', wrap=True, color=GREY); fvr = r; r += 1
    cell(sm, r, 1, 'Research target (12m) vs last close'); cell(sm, r, 3, f'=C{fvr}/Spot-1', fmt=PCT); r += 2

    # Summary: levels
    cell(sm, r, 1, 'Levels', bold=True, fill=SUB); cell(sm, r, 2, 'This model', bold=True, fill=SUB)
    cell(sm, r, 3, 'Published', bold=True, fill=SUB); cell(sm, r, 4, 'How', bold=True, fill=SUB); r += 1
    q = res['quantiles']
    rows = [
        ('NTM target (median)', q_formula(c_fin, 0.5), res['target'], 'Price with a 50% chance of closing above it on the horizon date'),
        ('Take-profit (75th percentile)', q_formula(c_fin, 0.25), res['take_profit'], '25% chance of closing above'),
        ('15th percentile', q_formula(c_fin, 0.85), q.get('q15'), '85% chance of closing above'),
    ]
    lv = {}
    for lab, f, pub, how in rows:
        cell(sm, r, 1, lab); cell(sm, r, 2, f.replace('=', '=', 1), fmt=USD, fill=OUT); cell(sm, r, 3, pub, color=BLUE, fmt=USD)
        cell(sm, r, 4, how, color=GREY, wrap=True); lv[lab] = r; r += 1
    q15r = lv['15th percentile']
    cell(sm, r, 1, 'Stop noise floor'); cell(sm, r, 2, '=Spot*EXP(-2*IV30*SQRT(21/252))', fmt=USD)
    cell(sm, r, 4, 'Two one-month implied moves below the last close; the stop is never tighter than this', color=GREY, wrap=True); nf = r; r += 1
    cell(sm, r, 1, 'Stop-loss', bold=True)
    cell(sm, r, 2, f'=MIN(IF(AND(Thesis_Break>0,Thesis_Break<Spot),MAX(B{q15r},Thesis_Break),B{q15r}),B{nf})', fmt=USD, bold=True, fill=OUT)
    cell(sm, r, 3, res['stop_loss'], color=BLUE, fmt=USD, bold=True)
    cell(sm, r, 4, f'Higher of the 15th percentile and the thesis break, capped at the noise floor (published basis: {res.get("stop_basis")})', color=GREY, wrap=True); r += 1
    tr = lv['NTM target (median)']
    cell(sm, r, 1, 'Target vs last close'); cell(sm, r, 2, f'=B{tr}/Spot-1', fmt=PCT); cell(sm, r, 3, f'=C{tr}/Spot-1', fmt=PCT); r += 1
    cell(sm, r, 1, 'P(close above last)'); cell(sm, r, 3, res['p_gain'], color=BLUE, fmt=PCT); r += 1
    cell(sm, r, 1, 'Expected price (mean)'); cell(sm, r, 3, res['expected'], color=BLUE, fmt=USD)
    cell(sm, r, 4, 'Mean of the blended distribution (combine.py integrates the full curve); above the median when right-skewed', color=GREY, wrap=True); r += 2
    cell(sm, r, 1, 'Sheets', bold=True, fill=SUB); cell(sm, r, 2, '', fill=SUB); cell(sm, r, 3, '', fill=SUB); cell(sm, r, 4, '', fill=SUB); r += 1
    for sh, what in [('Assumptions', 'The supervisor\'s table of every driver: each agent\'s value, reasoning and linked sources'),
                     ('Methods', 'Each agent\'s valuation methods, values and weights'),
                     ('Scenarios', 'Downside (10th pct) / base / upside (90th pct) price per agent'),
                     ('Distribution', 'The probability ladder and the blend, as formulas'),
                     ('Reasoning', 'Each agent\'s written reasoning, thesis break and what changed from round 1'),
                     ('Supervisor', 'Disagreements, fact-checks and changes'), ('Sources', 'Every cited source with its link')]:
        cell(sm, r, 1, sh, color=BLUE, link=f"#'{sh}'!A1"); cell(sm, r, 4, what, color=GREY); r += 1

    # ---------------- Assumptions ----------------
    asn = wb.create_sheet('Assumptions', 1)
    title(asn, f'{sym} key assumptions, aggregated by the supervisor',
          'Each driver of the valuation: the supervisor\'s reconciled value (grey rows), then what each agent assumed, why, and its sources. '
          'Blue values are the agents\' own numbers; blank means the agent did not use that driver.')
    A = ma.get('assumptions') or []
    header(asn, 4, ['ID', 'Assumption / who', 'Value', 'Reasoning', 'Source 1', 'Source 2', 'Source 3', 'All sources', 'Agents low', 'Agents high', 'Used in'],
           [6, 34, 13, 85, 28, 28, 28, 14, 12, 12, 22])

    def unit_fmt(u):
        u = (u or '').lower()
        if '%' in u: return PCT
        if u.strip() in ('x', 'multiple') or u.strip().endswith(' x'): return '0.0"x"'
        if '$' in u and 'share' in u: return USD
        if '$' in u: return '$#,##0.0;($#,##0.0);-'
        return NUM

    r = 5
    for a_ in A:
        fmt = unit_fmt(a_.get('unit')); hr = r
        rec = a_.get('reconciled') or {}
        cell(asn, r, 1, a_.get('id'), bold=True, fill=SUB)
        cell(asn, r, 2, f"{a_.get('name')} ({a_.get('unit')})" if a_.get('unit') else a_.get('name'), bold=True, fill=SUB, wrap=True)
        cell(asn, r, 3, num(rec.get('value')), color=BLUE, fmt=fmt, bold=True, fill=KEY)
        cell(asn, r, 4, ('Supervisor: ' + rec['reasoning']) if rec.get('reasoning') else '', wrap=True, fill=SUB)
        for k in range(3): asn.cell(row=r, column=5 + k).fill = SUB
        for k, u in enumerate([u for u in (rec.get('source_urls') or []) if str(u).startswith('http')][:3]):
            t = next((s_[2] for s_ in srows.values() if s_[1] == u), u)
            cell(asn, r, 5 + k, short(t, 45), color=BLUE, link=u, wrap=True, fill=SUB)
        cell(asn, r, 8, '', fill=SUB)
        cell(asn, r, 11, ', '.join(a_.get('used_in') or []), wrap=True, fill=SUB)
        r += 1; first = r
        for i, nm in enumerate(names):
            ag = (a_.get('agents') or {}).get(nm) or {}
            cell(asn, r, 1, a_.get('id'), color=GREY); cell(asn, r, 2, f'Agent {i+1}' + (f" · {ag['basis']}" if ag.get('basis') else ''))
            v = num(ag.get('value'))
            if v is not None: cell(asn, r, 3, v, color=BLUE, fmt=fmt)
            else: cell(asn, r, 3, short(ag.get('text') or 'not used', 30), color=GREY, wrap=True)
            txt = ag.get('reasoning') or ag.get('text') or ''
            cell(asn, r, 4, txt, wrap=True)
            ids = [x for x in (ag.get('source_ids') or []) if isinstance(x, int) and (nm, x) in srows]
            links = [(srows[(nm, x)][2], srows[(nm, x)][1]) for x in ids]
            adds = [x for x in (ag.get('added_sources') or []) if isinstance(x, dict) and x.get('url')]
            links += [('Supervisor: ' + (x.get('title') or x['url']), x['url']) for x in adds]
            for k, (t, u) in enumerate([l for l in links if str(l[1]).startswith('http')][:3]):
                cell(asn, r, 5 + k, short(t, 45), color=BLUE, link=u, wrap=True)
            aurl = [x['url'] for x in added]
            refs = [f'A{i+1}.{x}' for x in ids] + [f'X.{aurl.index(x["url"]) + 1}' for x in adds if x['url'] in aurl]
            if refs:
                first_row = srows[(nm, ids[0])][0] if ids else srows.get(('supervisor (added)', aurl.index(adds[0]['url']) + 1), (5,))[0]
                cell(asn, r, 8, ', '.join(refs), color=BLUE, link=f"#'Sources'!A{first_row}", wrap=True)
            elif v is not None:
                cell(asn, r, 8, 'no source found', color='C00000', bold=True)
            r += 1
        last = r - 1
        cell(asn, hr, 9, f'=IF(COUNT(C{first}:C{last})>0,MIN(C{first}:C{last}),"")', fmt=fmt, fill=SUB)
        cell(asn, hr, 10, f'=IF(COUNT(C{first}:C{last})>0,MAX(C{first}:C{last}),"")', fmt=fmt, fill=SUB)
        r += 1
    if not A:
        cell(asn, r, 1, 'The supervisor has not aggregated this round yet (model_assumptions.json); see Methods and Reasoning.', color=GREY)
    asn.freeze_panes = 'C5'

    # ---------------- Scenarios ----------------
    sc = wb.create_sheet('Scenarios')
    title(sc, f'{sym} cases by agent', 'Base case = median; downside = 10th percentile, upside = 90th percentile (no probabilities; older rounds\' bear/bull shown as downside/upside).')
    header(sc, 4, ['Agent', 'Downside (P10)', 'Base (median)', 'Upside (P90)', 'Downside story', 'Base story', 'Upside story'],
           [10, 13, 13, 13, 50, 50, 50])
    r = 5
    for i, a in enumerate(agents):
        s_ = a.get('scenarios') or {}
        cell(sc, r, 1, f'Agent {i+1}')
        for k, keys in enumerate([('downside', 'bear'), ('base',), ('upside', 'bull')]):
            d = next((s_[x] for x in keys if s_.get(x)), {})
            cell(sc, r, 2 + k, num(d.get('price')), color=BLUE, fmt=USD)
            cell(sc, r, 5 + k, d.get('story') or '', wrap=True)
        r += 1
    cell(sc, r, 1, 'Average', bold=True)
    for c in (2, 3, 4):
        cell(sc, r, c, f'=AVERAGE({CL(c)}5:{CL(c)}{r-1})', fmt=USD, bold=True)
    sc.freeze_panes = 'B5'

    # ---------------- Reasoning ----------------
    rs = wb.create_sheet('Reasoning')
    title(rs, f'{sym} each agent\'s reasoning', 'Final round, with round 1 for comparison. Revisions list what each agent changed after the supervisor\'s feedback.')
    header(rs, 4, ['Agent', 'Research target (12m) r1', 'Research target (12m) r2', 'Thesis break', 'Reasoning', 'Thesis break reason', 'Key uncertainties', 'Revisions after feedback'],
           [10, 12, 12, 12, 90, 50, 60, 70])
    r = 5
    for i, (nm, a) in enumerate(zip(names, agents)):
        a1 = r1.get(nm) or {}
        cell(rs, r, 1, f'Agent {i+1}')
        cell(rs, r, 2, num((a1.get('fair_value') or {}).get('value')), color=BLUE, fmt=USD)
        cell(rs, r, 3, num((a.get('fair_value') or {}).get('value')), color=BLUE, fmt=USD)
        cell(rs, r, 4, num(a.get('thesis_break_price')), color=BLUE, fmt=USD)
        cell(rs, r, 5, a.get('reasoning') or '', wrap=True)
        cell(rs, r, 6, a.get('thesis_break_reason') or '', wrap=True)
        cell(rs, r, 7, '\n'.join('• ' + str(x) for x in (a.get('key_uncertainties') or [])), wrap=True)
        rv = a.get('revisions') or []
        cell(rs, r, 8, '\n'.join(f'• {x.get("item", "")}: {x.get("from", "")} → {x.get("to", "")} ({x.get("why", "")})' if isinstance(x, dict) else f'• {x}' for x in rv), wrap=True)
        r += 1
    rs.freeze_panes = 'B5'

    # ---------------- Supervisor ----------------
    sv = wb.create_sheet('Supervisor')
    title(sv, f'{sym} supervisor', f'Confidence: {sup.get("confidence")}. {sup.get("confidence_reason") or ""}')
    sv.column_dimensions['A'].width = 30; sv.column_dimensions['B'].width = 60; sv.column_dimensions['C'].width = 60; sv.column_dimensions['D'].width = 50
    r = 4
    cell(sv, r, 1, 'Summary', bold=True); cell(sv, r, 2, sup.get('summary') or '', wrap=True); sv.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4); r += 1
    cell(sv, r, 1, 'Research target (12m)', bold=True); cell(sv, r, 2, (sfv or {}).get('reason') or '', wrap=True); sv.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4); r += 1
    cell(sv, r, 1, 'Key drivers', bold=True); cell(sv, r, 2, '\n'.join('• ' + str(x) for x in (sup.get('key_drivers') or [])), wrap=True)
    sv.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4); r += 2
    for lab, lst, keys in [('Disagreements (final round)', sup.get('disagreements') or [], ['topic', 'positions', 'resolution', 'evidence']),
                           ('Disagreements (round 1)', sup1.get('disagreements') or [], ['topic', 'positions', 'resolution', 'evidence']),
                           ('Changes the supervisor made', (sup.get('changes') or []) + (sup1.get('changes') or []), ['item', 'from', 'to', 'why'])]:
        if not lst: continue
        cell(sv, r, 1, lab, bold=True, size=12); r += 1
        for k, h in enumerate(keys): cell(sv, r, 1 + k, h.capitalize(), color='FFFFFF', bold=True, fill=HDR)
        r += 1
        for d in lst:
            for k, h in enumerate(keys): cell(sv, r, 1 + k, str(d.get(h, '') if isinstance(d, dict) else d), wrap=True)
            r += 1
        r += 1
    fb = sup1.get('feedback_for_agents') or []
    if fb:
        cell(sv, r, 1, 'Round-1 feedback sent back to the agents', bold=True, size=12); r += 1
        for k, h in enumerate(['Agent', 'What it missed', 'Revisit']): cell(sv, r, 1 + k, h, color='FFFFFF', bold=True, fill=HDR)
        r += 1
        for d in fb:
            cell(sv, r, 1, d.get('agent')); cell(sv, r, 2, '\n'.join('• ' + str(x) for x in (d.get('missed') or [])), wrap=True)
            cell(sv, r, 3, d.get('revisit') or '', wrap=True); r += 1
        r += 1
    ca = sup.get('citation_audit') or {}  # Kilian 2026-10-06: "everything must be cited accurately"
    if ca.get('corrections'):
        cell(sv, r, 1, f"Citation audit ({ca.get('date', '')})", bold=True, size=12); r += 1
        for x in ca['corrections']:
            cell(sv, r, 2, str(x), wrap=True); sv.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4); r += 1
        r += 1
    sz = sup.get('sizing') or {}
    if sz:
        cell(sv, r, 1, 'Sizing', bold=True, size=12); r += 1
        for k, v in sz.items():
            cell(sv, r, 1, k.replace('_', ' ')); cell(sv, r, 2, v if isinstance(v, (int, float)) else str(v), wrap=True); r += 1

    wb.move_sheet('Sources', offset=len(wb.sheetnames))
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)  # Excel recalculates on open
    out = os.path.join(work or rdir, f'model_{sym}.xlsx')
    wb.save(out)
    recalc(out)
    if work:  # built and recalculated on local disk, then one copy to the shared folder
        import shutil; final = os.path.join(rdir, f'model_{sym}.xlsx'); shutil.copyfile(out, final); return final
    return out


def recalc(path):
    """Store computed values too (for previewers), via LibreOffice Calc when it is installed; Excel recalculates anyway."""
    import subprocess, shutil
    rc = sorted(glob.glob('/root/.claude/skills/**/xlsx/scripts/recalc.py', recursive=True))
    if not rc or not shutil.which('soffice') or not glob.glob('/usr/lib/libreoffice/program/libsc*'):
        return
    r = subprocess.run(['python3', rc[0], path, '180'], capture_output=True, text=True)
    try:
        j = json.loads(r.stdout)
        if j.get('status') != 'success': print(os.path.basename(path), 'recalc:', r.stdout.strip()[:300])
    except Exception:
        print(os.path.basename(path), 'recalc failed:', (r.stdout or r.stderr)[:200])


def inputs_md5(date, sym):
    """Fingerprint of everything build() reads for this name, plus the code that shapes the workbook."""
    import hashlib
    h = hashlib.md5(); rdir = os.path.join(HERE, 'runs', date, sym)
    res = jl(os.path.join(rdir, 'result.json')) or {}
    adir = os.path.join(HERE, 'runs', res.get('research_as_of') or date, sym)
    files = [os.path.join(HERE, 'model_xlsx.py'), os.path.join(HERE, 'combine.py'), os.path.join(rdir, 'result.json'),
             os.path.join(adir, 'supervisor.json'), os.path.join(adir, 'round1', 'supervisor.json'), os.path.join(adir, 'model_assumptions.json')]
    files += sorted(glob.glob(os.path.join(adir, 'agent_*.json'))) + sorted(glob.glob(os.path.join(adir, 'round1', 'agent_*.json')))
    for f in files:
        h.update(f.encode())
        try: h.update(open(f, 'rb').read())
        except OSError: h.update(b'-')
    return h.hexdigest()


def _one(a):
    date, s, force = a
    import tempfile, io, contextlib
    rdir = os.path.join(HERE, 'runs', date, s); fp_path = os.path.join(rdir, f'model_{s}.inputs.md5')
    fp = inputs_md5(date, s)
    if not force and os.path.exists(os.path.join(rdir, f'model_{s}.xlsx')) and (jl(fp_path) or {}).get('md5') == fp:
        return s, os.path.join(rdir, f'model_{s}.xlsx'), '', True
    buf = io.StringIO()
    with tempfile.TemporaryDirectory(prefix=f'model_{s}_') as work, contextlib.redirect_stdout(buf):
        p = build(date, s, work)
    if p: json.dump(dict(md5=fp), open(fp_path, 'w'))
    return s, p, buf.getvalue(), False


if __name__ == '__main__':
    args = sys.argv[1:]; jobs = 4; force = '--force' in args
    if '--jobs' in args: i = args.index('--jobs'); jobs = int(args[i + 1]); del args[i:i + 2]
    args = [x for x in args if x != '--force']
    date = args[0]
    syms = args[1:] or sorted(os.path.basename(p) for p in glob.glob(os.path.join(HERE, 'runs', date, '*')) if os.path.isdir(p))
    import concurrent.futures as cf
    with cf.ProcessPoolExecutor(max(1, min(jobs, len(syms) or 1))) as ex:
        kept = 0
        for s, p, out, same in ex.map(_one, [(date, s, force) for s in syms]):
            if out: print(out, end='')
            if p: print(p + (' (inputs unchanged, kept)' if same else ''))
            kept += same
    print(f'{len(syms) - kept} built, {kept} kept (inputs unchanged)')
